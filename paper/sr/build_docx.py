"""Builds the Word version of the paper in the layout of the reference document (a Scientific Reports article):
single text column with a wide left margin, Corbel headings, serif body, light-blue running header, numbered
superscript citations, grey-header tables with captions below, figures with captions below, and the journal's
end-matter sections. Publisher marks (logo, DOI, volume, licence text) are intentionally not reproduced.

    python build_docx.py            ->  Verify_before_you_write_Jerath_Jagadeesan.docx
"""
import os, re, sys, copy
from docx import Document
from docx.shared import Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import tex2blocks as T
from refs_data import REFS

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'Verify_before_you_write_Jerath_Jagadeesan.docx')

# ------------------------------------------------------------------------------------------------ page geometry (pt)
PAGE_W, PAGE_H = 595.28, 782.36
LEFT, RIGHT, TOP, BOTTOM = 155.9, 39.1, 46.0, 40.0
TEXT_W = PAGE_W - LEFT - RIGHT          # 400.3
WIDE_W = 508.0                          # tables / diagrams that use the full page width
WIDE_INDENT = -(LEFT - 46.9)            # they start at x = 46.9 pt
BODY_FONT, HEAD_FONT, MONO_FONT = 'Times New Roman', 'Corbel', 'Consolas'
BLUE = RGBColor(0x00, 0x00, 0xFF)
BAND = 'CEDDE3'
HEAD_FILL = 'A8A9AC'

TITLE = 'Verify before you write: workbook-grounded verification for LLM spreadsheet agents'
AUTHORS = ['Ansh Jerath', 'Jagadeesan S']
AFFIL = ('School of Computer Science and Engineering and Information Systems (SCORE), Vellore Institute of Technology (VIT), '
         'Vellore, Tamil Nadu, India.')
EMAILS = 'ansh.jerath2023@vitstudent.ac.in; s.jagadeesan@vit.ac.in'
KEYWORDS = 'Spreadsheet formulas, Large language models, LLM agents, Verification, Benchmark, Retrieval, SSRF'

NOTE_TITLES = ['Benchmark details', 'Verifier details', 'Prompts', 'Additional results',
               'Checking the labels in Google Sheets', 'Independent annotation of real formulas']
NOTE_LABELS = ['app:bench', 'app:verifier', 'app:prompts', 'app:results', 'app:gsheets', 'app:annotation']
FIG_WIDTH = {'fig_arch': 508, 'fig_anatomy': 508, 'fig_bench': 508, 'fig_rw_overlap': 312,
             # tall single-column figures are narrowed so that they fit in the space left at the bottom of a page more often
             'fig_e1_classes': 250, 'fig_e2_budget': 262, 'fig_e5_templates': 280, 'fig_e5_repair_fate': 290}
FIG_DEFAULT = 300


# -------------------------------------------------------------------------------------------------- content helpers
def tex(s):
    return T.parse_inline(T.preprocess(s))


def P(s):
    return ('p', tex(s))


def H1(title, label=None):
    return ('h1', [T.run(title)], label)


def H2(title, label=None):
    return ('h2', [T.run(title)], label)


def text_of(runs):
    return T.runs_text(runs)


def L(rel):
    return T.load_file(rel)


def split_by(blocks, kind):
    """Split a block list at headings of the given kind: returns (head, [(heading, blocks), ...])."""
    head, parts, cur = [], [], None
    for b in blocks:
        if b[0] == kind:
            cur = (b, [])
            parts.append(cur)
        elif cur is None:
            head.append(b)
        else:
            cur[1].append(b)
    return head, parts


# ------------------------------------------------------------------------------------------------- Table 1 (related)
def cell(s, **kw):
    # a bare % would start a LaTeX comment, so literal percent signs in the cell text are escaped first
    return {'runs': tex(s.replace('%', r'\%')), 'span': 1, 'vspan': 1, 'vcont': False, **kw}


def related_table():
    hdr = ['S no', 'Paper', 'Dataset or task', 'Objective', 'Result', 'Gap relative to this work']
    rows = [
        ['1', r'Zhao et al.\cite{zhao2024nl2formula}', 'NL2Formula', 'Table-grounded formula generation from natural-language queries, evaluated by execution', '–', 'Judges the final formula; no check before the write'],
        ['2', r'Joshi et al.\cite{joshi2024flame}', 'Excel formulas', 'FLAME, a small language model trained on spreadsheet formulas', 'Competes with larger models', 'Final-answer accuracy only'],
        ['3', r'Li et al.\cite{li2023sheetcopilot}', 'Spreadsheet manipulation tasks', 'SheetCopilot maps requests to atomic actions', '–', 'Task success, not verification of a candidate'],
        ['4', r'Ma et al.\cite{ma2024spreadsheetbench}', 'SpreadsheetBench (912 real forum questions)', 'Benchmark of real-world spreadsheet manipulation', '–', 'Measures outcomes; no verifier under test'],
        ['5', r'Dong et al.\cite{dong2024spreadsheetllm}', '–', 'SpreadsheetLLM compresses a sheet to fit a prompt', '–', 'Representation of the sheet, not verification'],
        ['6', r'Chen et al.\cite{chen2025sheetmind}', '–', 'SheetMind studies what determines how an agent fails', '–', 'Analyses failures; does not stop them'],
        ['7', r'Thorne\cite{thorne2025flare}', '–', 'FLARE targets spreadsheet auditing', '–', 'Audits existing sheets, not model output before the write'],
        ['8', r'Singha et al.\cite{singha2025forepbench}', 'Synthetic repair data', 'Benchmark data for repairing formulas that already raise runtime errors, with LLM and execution judging', '–', 'Repairs visible errors; silent faults not covered'],
        ['9', r'Singh et al.\cite{singh2025validating}', 'Synthetic formula descriptions', 'LLM-based validators clean fine-tuning data for formula generation', 'Validation improves performance over raw data across four models', 'Validates training data, not candidate formulas'],
        ['10', r'Tian et al.\cite{tian2025sheetpedia}', 'Sheetpedia (over 290,000 spreadsheets)', 'Corpus for spreadsheet intelligence and LLM fine-tuning (NL2SR, NL2Formula)', 'Up to 97.5% (NL2SR) and 71.7% (NL2Formula) accuracy after fine-tuning', 'No cached results; used here as a source of real formulas'],
        ['11', r'Ren et al.\cite{ren2026spreadsheetagent}', 'Spreadsheet Bench', 'SpreadsheetAgent, a two-stage multi-agent framework with a module that validates extracted structure', '38.16% with GPT-OSS-120B against 35.27% for the ChatGPT Agent baseline', 'Verifies extraction of a sheet, not a candidate formula'],
        ['12', r'Zhu et al.\cite{zhu2026spreadsheetbench2}', 'SpreadsheetBench 2 (321 workflow tasks)', 'End-to-end business workflows (generation, debugging, visualization) with eight frontier models', 'Best model 34.89% overall; best debugging accuracy 12%', 'Workflow outcome; does not isolate pre-write checking'],
    ]
    grid = [[cell(h) for h in hdr]] + [[cell(c) for c in r] for r in rows]
    for r in grid:
        for i, c in enumerate(r):
            c['col'] = i
    cols = [{'a': 'l', 'w': 0.9}, {'a': 'l', 'w': 2.2}, {'a': 'l', 'w': 3.0}, {'a': 'l', 'w': 4.6}, {'a': 'l', 'w': 3.4}, {'a': 'l', 'w': 3.4}]
    return ('table', {'cols': cols, 'rows': grid, 'header_rows': 1, 'wide': True, 'label': 'tab:related',
                      'caption': tex('Comparative review of approaches to spreadsheet formula generation, validation and repair.')})


# ---------------------------------------------------------------------------------------------------- Algorithm 1
ALGO = [
    'BEGIN',
    'Step 1: Parse the candidate formula',
    '  TOKENIZE formula; PARSE it into an abstract syntax tree by recursive descent',
    '  IF parsing fails THEN errors ← syntax error (structural layer)',
    'Step 2: Read the open workbook',
    '  sheet_names, used_extent(sheet), headers, named_ranges ← workbook',
    '  column_profile(column) ← type, distinct values, sample (at most 500 rows per referenced column)',
    'Step 3: Run the six layers on the tree',
    '  STRUCTURAL: quotes, parentheses, injection patterns',
    '  SYMBOLS: every sheet, function, name and QUERY column exists (a near-miss function name is an error)',
    '  BOUNDS: every range lies inside the extent of the sheet it points at',
    '  SHAPE: argument counts, lookup index, unequal range sizes',
    '  GROUNDING: numeric aggregate over a text column; criterion absent from its column ← suspicious',
    '  CIRCULARITY: target cell inside a referenced range or on a dependency chain',
    'Step 4: Decide',
    '  IF errors is empty THEN ACCEPT the formula and write it',
    '  ELSE REJECT',
    'Step 5: Repair (at most two retries)',
    '  FOR attempt ← 1 TO 2 DO',
    '    APPEND formula, errors, warnings and hints to the conversation',
    '    temperature ← max(0, 0.2 − 0.1 × attempt)',
    '    candidate ← AGENT(conversation); REPEAT Steps 1 to 4',
    '  IF the formula is still rejected THEN do not write it; return the diagnostics',
    'END',
]


# ------------------------------------------------------------------------------------------------------ assembly
def assemble():
    doc = []

    # ---- Introduction (untitled, as in the reference) -------------------------------------------------------
    intro = L('sections/intro.tex')[1:]
    doc += intro
    doc.append(P(r'The rest of the paper is structured as follows. \xref{Section}{sec:related} reviews the literature and compares the approaches closest to this work. '
                 r'\xref{Section}{sec:limitations} is devoted to the possible threats to validity that can influence the results. '
                 r'\xref{Section}{sec:gap} outlines the research gap and justifies the significance of the proposed approach. '
                 r'\xref{Section}{sec:method} describes the general methodology, whereas the proposed system is presented in \xref{Section}{sec:system}. '
                 r'\xref{Section}{sec:impl} outlines implementation aspects and the experimental setup, and \xref{Section}{sec:results} reports and discusses the results. '
                 r'\xref{Section}{sec:conclusion} finalizes the paper with a conclusion on the key findings and future research directions.'))

    # ---- Related work -------------------------------------------------------------------------------------
    doc.append(H1('Related work', 'sec:related'))
    doc += T.reorder_floats(L('sections/related.tex')[1:])
    doc.append(P(r'\xref{Table}{tab:related} compares the approaches reviewed above with the present study.'))
    doc.append(related_table())

    # ---- Threats to validity (the first part of the limitations section) -----------------------------------
    lim = L('sections/limitations.tex')
    head, parts = split_by(lim, 'h1')
    threats = parts[0][1]
    ethics_repro = parts[1][1]
    doc.append(H1('Threats to validity', 'sec:limitations'))
    doc += T.reorder_floats(threats)

    # ---- Research gap -------------------------------------------------------------------------------------
    doc.append(H1('Research gap', 'sec:gap'))
    doc.append(P(r'Language models can now write spreadsheet formulas, but the step between a model’s output and a change to the user’s workbook is still weakly studied. '
                 r'To our knowledge, the following gaps remain open.'))
    gap_items = [
        r'Work on spreadsheet language models mostly assesses the final answer, through execution accuracy on formula generation or task success on end-to-end workflows\cite{zhao2024nl2formula,ma2024spreadsheetbench,zhu2026spreadsheetbench2}, and not the checks that run on a candidate formula before it is committed.',
        r'Where verification appears, it targets other objects: it validates synthetic training data\cite{singh2025validating} or the extraction of a sheet’s structure\cite{ren2026spreadsheetagent}, or it repairs formulas that already raise errors\cite{singha2025forepbench}. Deterministic verification of a generated formula against workbook state, scored against labels from an independent engine and with an explicit distinction between loud failures (an error value) and silent failures (a plausible wrong value), has not, to our knowledge, been reported.',
        r'A benchmark written together with the checker it evaluates cannot show which faults the authors did not think of, and corpora of real formulas such as Sheetpedia\cite{tian2025sheetpedia} carry no cached results, so no ground truth for real defects is available.',
        r'What verification-guided repair does to the visibility of failures, that is, whether it removes wrong formulas or turns visible errors into plausible wrong values, is not captured by accuracy and has not been measured.',
        r'How much a deterministic verifier helps depends on the model that writes the formulas, and small local models and larger hosted models behave very differently.',
    ]
    doc.append(('list', [tex(s) for s in gap_items]))

    # ---- Methodology --------------------------------------------------------------------------------------
    meth = L('sections/method.tex')
    mhead, mparts = split_by(meth, 'h1')
    mbody = mparts[0][1]
    mintro, msubs = split_by(mbody, 'h2')
    bench_fig = None
    sub_blocks = {}
    for h, bl in msubs:
        title = text_of(h[1])
        keep = []
        for b in bl:
            if b[0] == 'figure' and b[1]['label'] == 'fig:bench':
                bench_fig = b
            else:
                keep.append(b)
        sub_blocks[title] = (h, keep)
    doc.append(H1('Methodology', 'sec:method'))
    doc.append(P(r'The research methodology of this work is organised around one question, namely what deterministic checks that read only the open workbook can decide about a candidate formula before it is written, and four research questions (RQ1 to RQ4) that make it testable. '
                 r'The method has five stages: construction of a verifier, construction of a benchmark with independent labels, scoring of real formulas that people wrote, an end-to-end study with real language models, and statistics that respect how the data are clustered. '
                 r'The stages are designed so that each question is answered with evidence that does not depend on the component under test.'))
    doc += mintro
    doc.append(P(r'In the first stage, a candidate formula proposed by the add-on’s formula agent is parsed and checked in six layers against the open workbook (\xref{Section}{sec:system}). '
                 r'In the second stage, SheetFault corrupts correct formulas with \nBenchClasses{} fault operators and labels every result loud or silent with an independent spreadsheet engine, which gives recall and false-positive rates for any verifier (\xref{Figure}{fig:bench}; \xref{Section}{sec:benchmark}). '
                 r'In the third stage, three disjoint sets of 300 Sheetpedia workbooks provide real formulas, the last two scored once with the verifier frozen in advance (\xref{Section}{sec:rwmethod}). '
                 r'In the fourth stage, the production formula agent runs end to end with two local and three hosted language models under five repair strategies and three write policies, and each final cell is scored by execution (\xref{Section}{sec:e5method}). '
                 r'In the fifth stage, intervals respect the clustering of formulas within workbooks and p-values are corrected for multiple comparisons (\xref{Section}{sec:statsmethod}).'))
    doc.append(bench_fig)

    # ---- Proposed system ----------------------------------------------------------------------------------
    sysb = L('sections/system.tex')
    doc.append(H1('Proposed system', 'sec:system'))
    system_blocks = [b for b in sysb[1:]]
    # system.tex: p, figure(arch), h2 gc, list.., p, figure(anatomy), h2 levels, p, table, h2 retrieval, p, p
    sys_pre, sys_parts = split_by(system_blocks, 'h2')
    doc += T.reorder_floats(sys_pre)
    for h, bl in sys_parts:
        title = text_of(h[1])
        if title.startswith('Retrieval'):
            doc.append(H2('Strengths of the proposed system', None))
            doc.append(P(r'Four properties distinguish the verifier from the checks it replaces.'))
            doc.append(('list', [tex(s) for s in [
                r'\textbf{Deterministic and cheap.} The verifier makes no model call: it parses the candidate and reads the open workbook, so the same formula in the same workbook always receives the same verdict, and a check costs a number of \code{SpreadsheetApp} calls that grows with the number of sheets but not with the number of rows (\xref{Section}{sec:e3}).',
                r'\textbf{Specific feedback.} A rejection names what the workbook actually holds, such as the column letter of a header, the real values of a column or the real sheet names, which is the information a model needs to repair the formula (\xref{Figure}{fig:anatomy}).',
                r'\textbf{Blocks only what the workbook proves.} A finding that the workbook proves impossible, or a near-miss of something real, blocks the write; a plausible but probably wrong finding is demoted to a warning.',
                r'\textbf{Independent of the model.} It checks the candidate and not its author, so the same verifier serves local and hosted models; how much it helps depends on how many errors a model makes that the workbook can see (\xref{Section}{sec:e5}).',
            ]]))
            doc.append(P(r'Its limit is stated in \xref{Section}{sec:levels}: it cannot decide whether a valid formula answers the question the user asked.'))
            # the algorithm goes before the retrieval subsection (it summarises the verifier described above)
            doc.append(H2('Pseudo code for the proposed system', None))
            doc.append(P(r'\xref{Algorithm}{alg:gc} summarises the verification and repair loop in pseudo code.'))
            doc.append(('algorithm', {'lines': ALGO, 'label': 'alg:gc', 'caption': tex('Verification and repair of a candidate spreadsheet formula (GroundCheck).')}))
        doc.append(h)
        doc += T.reorder_floats(bl)

    # ---- Implementation -----------------------------------------------------------------------------------
    doc.append(H1('Implementation', 'sec:impl'))
    doc.append(P(r'The experiments run the add-on’s own source files inside a sandbox that stubs only the Apps Script globals. The evaluation harness is written in Node.js, uses HyperFormula\cite{hyperformula} as the independent spreadsheet engine, and every generator is seeded. '
                 r'This section describes the benchmark, the real-formula sets, the end-to-end study, the statistics, and the software and data needed to reproduce them.'))
    rename = {'SheetFault': ('SheetFault benchmark', 'sec:benchmark'), 'Real-world formulas': ('Real-world formula sets', 'sec:rwmethod'),
              'End-to-end study': ('End-to-end study', 'sec:e5method'), 'Statistics and metrics': ('Statistics and metrics', 'sec:statsmethod'),
              'Retrieval, cost and security': ('Retrieval, cost and security', None)}
    for title, (h, bl) in sub_blocks.items():
        new_title, lab = rename[title]
        doc.append(H2(new_title, lab or h[2]))
        doc += T.reorder_floats(bl)
    repro = [b for b in ethics_repro if b[0] != 'h3' or text_of(b[1]) == 'Reproducibility']
    # keep only the Reproducibility paragraph for this section
    ri = next(i for i, b in enumerate(ethics_repro) if b[0] == 'h3' and text_of(b[1]) == 'Reproducibility')
    ei = next(i for i, b in enumerate(ethics_repro) if b[0] == 'h3' and text_of(b[1]) == 'Ethics')
    doc.append(H2('Reproducibility', None))
    doc.append(('p', ethics_repro[ri + 1][1]))
    ethics_block = ethics_repro[ei + 1]

    # ---- Results and analysis -----------------------------------------------------------------------------
    res = L('sections/results.tex')
    doc.append(H1('Results and analysis', 'sec:results'))
    for b in T.reorder_floats(res[1:]):
        if b[0] == 'h2' and b[2] == 'sec:e3':
            b = ('h2', [T.run('Cost and the fetch boundary')], 'sec:e3')
        doc.append(b)
    doc.append(H2('Discussion', 'sec:discussion'))
    doc += T.reorder_floats(L('sections/discussion.tex')[1:])

    # ---- Conclusion ---------------------------------------------------------------------------------------
    doc.append(H1('Conclusion', 'sec:conclusion'))
    doc.append(P(r'This work asked what deterministic checks that read only the open workbook can decide about a candidate spreadsheet formula before it is written, and what they cannot. '
                 r'We built GroundCheck, a parser-based verifier with six workbook-grounded layers and structured repair feedback, and SheetFault, a seeded benchmark of \nBenchTotal{} formulas whose loud and silent faults are labelled by an independent spreadsheet engine, and we evaluated both on the benchmark, on real formulas, and end to end with two local and three hosted language models.'))
    doc.append(P(r'The benchmark shows a large gap: GroundCheck blocks \eonevtworejectR\% of faults against \eonevonerejectR\% for the verifier shipped in the add-on, at \eonevtworejectFPR\% false positives, and it is ahead in \clWbBetter{} of \clWbN{} workbooks. '
                 r'The benchmark was written with the verifier, however, and the real formulas tell a more modest story. On three disjoint sets of about 5{,}000 formulas the two verifiers reject a similar share (\rwcTVtwoReject\%, \rwcFVtwoReject\% and \rwcGVtwoReject\% for GroundCheck), the difference between them is within workbook-clustered noise, and the frozen confirmation run exposed a class of real defects, deleted references and \code{INDIRECT} calls naming missing sheets, that neither the benchmark nor the first version of GroundCheck covered. '
                 r'Two rules added in response change no earlier verdict and were confirmed on a second untouched set.'))
    doc.append(P(r'End to end, GroundCheck feedback improves an 8B model by \sdADvtwofbresample{} points over plain resampling, in each of \sdASeeds{} seeds, and a 3B model by \sdBDvtwofbresample{}, while hosted models leave the verifier idle because they make almost no errors the workbook can see. '
                 r'Repair has a cost that accuracy hides: it raises the number of plausible-but-wrong formulas written, from \hmAAsIsSilent{} to \hmARepVtwoSilent{} of \hmAN{} tasks for the 8B model, whereas holding the rejected formula and showing the diagnostics does not. An execution check removes some visible errors but not silent ones. '
                 r'A workbook can decide syntax, names, extent, shape and dependencies; it cannot decide whether a valid formula answers the user’s question.'))
    doc.append(P(r'Future work follows the threats to validity. The benchmark and the real-formula sets should be extended with independent human labels and with evaluation in Google Sheets itself, for which a harness is provided. '
                 r'Larger and harder multi-step tasks and further model families, including closed models that we could not run, are needed to learn when verification matters for strong models. '
                 r'The benchmark should cover dates, arrays, text manipulation and structured references. A user study should measure whether people notice the plausible-but-wrong values that repair can create, and whether showing the computed value next to the request prevents them. '
                 r'The benchmark, the verifier, the frozen real-formula protocol and the measure of unsafe repair are offered as a foundation for such studies.'))

    # ---- end matter ---------------------------------------------------------------------------------------
    doc.append(H1('Data availability', None))
    doc.append(P(r'The SheetFault generators, the evaluation harness, the raw model trajectories and the analysis scripts that accompany this study are in the project repository (https://github.com/darkhorse0204/sheets-ai-pipeline) and are available from the corresponding authors on request. '
                 r'The Sheetpedia workbooks are publicly available from their original release (https://huggingface.co/datasets/tianzl66/Sheetpedia\_xlsx) under CC BY-SA 4.0 and are not redistributed here.'))
    doc.append(('marker', 'references'))
    doc.append(H1('Author contributions', None))
    doc.append(P(r'Ansh Jerath: Conceptualization, Methodology, Software, Investigation, Formal analysis, Data curation, Visualization, Writing – original draft. '
                 r'Jagadeesan S: Supervision, Project administration, Writing – review \& editing.'))
    doc.append(H1('Funding', None))
    doc.append(P(r'The authors received no specific funding for this work.'))
    doc.append(H1('Declarations', None))
    doc.append(H1('Competing interests', None))
    doc.append(P(r'The authors declare no competing interests.'))
    doc.append(H1('Ethical considerations', None))
    doc.append(('p', ethics_block[1]))
    doc.append(H1('Additional information', None))
    doc.append(P(r'\textbf{Correspondence} and requests for materials should be addressed to A.J. or J.S.'))

    # ---- supplementary information ------------------------------------------------------------------------
    doc.append(('marker', 'supp'))
    doc.append(H1('Supplementary information', None))
    doc.append(P('The supplementary notes give the benchmark details, the verifier details, the prompts, additional results, and the protocols for the two checks that need a person or a Google account.'))
    app = L('sections/appendix.tex')
    ahead, aparts = split_by(app, 'h1')
    for n, (h, bl) in enumerate(aparts):
        doc.append(('h2', [T.run(f'Supplementary Note {n + 1}. {NOTE_TITLES[n]}')], NOTE_LABELS[n]))
        doc += T.reorder_floats(bl)
    return doc


# --------------------------------------------------------------------------------------------- numbering / refs
class Registry:
    def __init__(self):
        self.labels = {}        # label -> (kind, display number)
        self.cite_no = {}       # key -> n
        self.cite_order = []
        self.note_no = {l: i + 1 for i, l in enumerate(NOTE_LABELS)}
        self.titles = {}        # section label -> heading text


def number(doc):
    reg = Registry()
    cnt = {'Table': 0, 'Figure': 0, 'STable': 0, 'SFigure': 0, 'Algorithm': 0}
    supp = False

    def scan(runs):
        for r in runs:
            for k in r.get('cite', []):
                if k not in reg.cite_no:
                    reg.cite_order.append(k)
                    reg.cite_no[k] = len(reg.cite_order)

    for b in doc:
        k = b[0]
        if k == 'marker' and b[1] == 'supp':
            supp = True
        if k in ('h1', 'h2') and b[2]:
            reg.titles[b[2]] = text_of(b[1])
        if k in ('p',):
            scan(b[1])
        elif k == 'list':
            for it in b[1]:
                scan(it)
        elif k == 'table':
            t = b[1]
            for row in t['rows']:
                for c in row:
                    scan(c['runs'])
            scan(t['caption'])
            key = 'STable' if supp else 'Table'
            cnt[key] += 1
            if t['label']:
                reg.labels[t['label']] = ('Table', ('S%d' if supp else '%d') % cnt[key], supp)
        elif k == 'figure':
            f = b[1]
            scan(f['caption'])
            key = 'SFigure' if supp else 'Figure'
            cnt[key] += 1
            if f['label']:
                reg.labels[f['label']] = ('Figure', ('S%d' if supp else '%d') % cnt[key], supp)
        elif k == 'algorithm':
            cnt['Algorithm'] += 1
            reg.labels[b[1]['label']] = ('Algorithm', str(cnt['Algorithm']), False)
    # citation numbers must follow the order of appearance in the final flow, including captions placed after tables
    return reg


# -------------------------------------------------------------------------------------------------- docx helpers
def set_run_font(run, name=None, size=None, bold=None, italic=None, color=None, raise_pt=None):
    f = run.font
    if raise_pt:
        # the reference raises smaller digits by shifting the baseline (it does not use the superscript attribute)
        pos = OxmlElement('w:position'); pos.set(qn('w:val'), str(int(raise_pt * 2)))
        run._r.get_or_add_rPr().append(pos)
    if name:
        f.name = name
        rpr = run._r.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(a), name)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def style_fonts(style, name, size, bold=False, italic=False):
    f = style.font
    f.name = name
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)


def make_style(doc, name, base='Normal', font=BODY_FONT, size=9, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
               line=10, before=0, after=0, first=0, left=0, keep_next=False, keep_lines=False, outline=None, color=None, nohyph=False):
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles[base]
    style_fonts(st, font, size, bold, italic)
    if color is not None:
        st.font.color.rgb = color
    pf = st.paragraph_format
    pf.alignment = align
    if line:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line)
    else:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Pt(first)
    pf.left_indent = Pt(left)
    pf.keep_with_next = keep_next
    pf.keep_together = keep_lines
    pf.widow_control = True
    if outline is not None:
        ppr = st.element.get_or_add_pPr()
        o = OxmlElement('w:outlineLvl'); o.set(qn('w:val'), str(outline)); ppr.append(o)
    if nohyph:
        ppr = st.element.get_or_add_pPr()
        ppr.append(OxmlElement('w:suppressAutoHyphens'))
    return st


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(PAGE_W), Pt(PAGE_H)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Pt(LEFT), Pt(RIGHT), Pt(TOP), Pt(BOTTOM)
    sec.header_distance, sec.footer_distance = Pt(0), Pt(13)

    # document defaults: no theme fonts
    styles_el = doc.styles.element
    dd = styles_el.find(qn('w:docDefaults'))
    if dd is not None:
        rf = dd.find('.//' + qn('w:rFonts'))
        if rf is not None:
            for a in list(rf.attrib):
                del rf.attrib[a]
            for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
                rf.set(qn(a), BODY_FONT)
        lang = dd.find('.//' + qn('w:lang'))
        if lang is not None:
            lang.set(qn('w:val'), 'en-US')
    normal = doc.styles['Normal']
    style_fonts(normal, BODY_FONT, 9)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(10)
    pf.space_before = pf.space_after = Pt(0)
    pf.first_line_indent = Pt(12)
    pf.widow_control = True

    make_style(doc, 'SR Body first', first=0)
    make_style(doc, 'SR Title', font=HEAD_FONT, size=26, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, line=30, first=0)
    make_style(doc, 'SR Authors', font=HEAD_FONT, size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, line=12, before=8, after=9, first=0)
    make_style(doc, 'SR Abstract', font=HEAD_FONT, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, line=11, after=8, first=0, nohyph=True)
    make_style(doc, 'SR Keywords', size=9, align=WD_ALIGN_PARAGRAPH.LEFT, line=11, after=9, first=0)
    make_style(doc, 'SR Affil', font=HEAD_FONT, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT, line=10, first=0)
    make_style(doc, 'SR Heading 1', font=HEAD_FONT, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, line=13, before=11, after=2, first=0, keep_next=True, outline=0)
    make_style(doc, 'SR Heading 2', font=HEAD_FONT, size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, line=12, before=9, after=1.5, first=0, keep_next=True, outline=1)
    make_style(doc, 'SR Bullet', first=-9.9, left=9.9, before=0, after=0)
    make_style(doc, 'SR Quote', font=MONO_FONT, size=7.5, align=WD_ALIGN_PARAGRAPH.LEFT, line=9.5, first=0, left=14)
    make_style(doc, 'SR Caption', size=9, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=10, before=5, after=12, first=0, keep_lines=True)
    make_style(doc, 'SR Figure', align=WD_ALIGN_PARAGRAPH.CENTER, line=None, before=6, after=0, first=0, keep_next=True)
    make_style(doc, 'SR Table text', size=7, align=WD_ALIGN_PARAGRAPH.LEFT, line=8.6, first=0, nohyph=True)
    make_style(doc, 'SR Table gap', size=2, align=WD_ALIGN_PARAGRAPH.LEFT, line=4, first=0, keep_next=True)
    make_style(doc, 'SR Reference', size=7.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=8.4, first=-12.8, left=12.8)
    make_style(doc, 'SR Algorithm', font='Cambria', size=8.2, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, line=10, first=0, left=22)
    make_style(doc, 'SR Header', font=HEAD_FONT, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, line=28, first=0)
    make_style(doc, 'SR Footer', font=HEAD_FONT, size=8, align=WD_ALIGN_PARAGRAPH.LEFT, line=10, first=0)

    # automatic hyphenation (the reference hyphenates)
    st = doc.settings.element
    after = st.find(qn('w:defaultTabStop'))
    for tag, val in (('w:autoHyphenation', None), ('w:consecutiveHyphenLimit', '3'), ('w:hyphenationZone', '300')):
        el = OxmlElement(tag)
        if val:
            el.set(qn('w:val'), val)
        if after is not None:
            after.addnext(el)
            after = el
        else:
            st.append(el)

    # the first page carries only a thin band (the reference puts its logo there); later pages carry the running title
    sec.different_first_page_header_footer = True
    build_header(doc, sec.header, TITLE, 28)
    build_header(doc, sec.first_page_header, '', 9)
    build_footer(doc, sec.footer)
    build_footer(doc, sec.first_page_footer)


def build_header(doc, header, text, height):
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.style = doc.styles['SR Header']
    ppr = hp._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), BAND)
    ppr.append(shd)
    hp.paragraph_format.line_spacing = Pt(height)
    # the band runs from page edge to page edge; the running title starts at the left edge of the footer rule
    hp.paragraph_format.left_indent = Pt(-LEFT)
    hp.paragraph_format.first_line_indent = Pt(41.1)
    hp.paragraph_format.right_indent = Pt(-RIGHT)
    if text:
        set_run_font(hp.add_run(text), HEAD_FONT, 10)


def build_footer(doc, footer):
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.style = doc.styles['SR Footer']
    fp.paragraph_format.left_indent = Pt(-(LEFT - 41.1))
    fp.paragraph_format.right_indent = Pt(556.2 - 554.2)
    ppr = fp._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    for a, v in (('w:val', 'single'), ('w:sz', '2'), ('w:space', '3'), ('w:color', '000000')):
        top.set(qn(a), v)
    bdr.append(top)
    ppr.append(bdr)
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab'); tab.set(qn('w:val'), 'right'); tab.set(qn('w:pos'), str(int((554.2 - LEFT) * 20)))
    tabs.append(tab)
    ppr.append(tabs)
    r1 = fp.add_run('Manuscript'); set_run_font(r1, HEAD_FONT, 8, bold=True)
    r2 = fp.add_run('  |  Jerath & Jagadeesan'); set_run_font(r2, HEAD_FONT, 8)
    r3 = fp.add_run('\t'); set_run_font(r3, HEAD_FONT, 8)
    r4 = fp.add_run(); set_run_font(r4, HEAD_FONT, 10, color=RGBColor(0x59, 0x59, 0x59))
    for typ, txt in (('begin', None), (None, ' PAGE '), ('separate', None), ('t', '1'), ('end', None)):
        if typ in ('begin', 'separate', 'end'):
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), typ); r4._r.append(fc)
        elif typ == 't':
            t_ = OxmlElement('w:t'); t_.text = txt; r4._r.append(t_)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; r4._r.append(it)


# ---------------------------------------------------------------------------------------------------- runs
class Ctx:
    def __init__(self, reg):
        self.reg = reg


def cite_text(keys, reg):
    nums = sorted({reg.cite_no[k] for k in keys})
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            out.append(f'{nums[i]}–{nums[j]}')
        else:
            out.extend(str(x) for x in nums[i:j + 1])
        i = j + 1
    return ','.join(out)


def xref_runs(word, label, sofar, reg):
    blue = {'blue': True}
    if label.startswith('sec:'):
        return [{'t': 'Section “'}, {'t': reg.titles[label], **blue}, {'t': '”'}]
    if label.startswith('app:'):
        return [{'t': f'Supplementary Note {reg.note_no[label]}', **blue}]
    kind, n, supp = reg.labels[label]
    if kind == 'Table':
        return [{'t': ('Supplementary Table ' if supp else 'Table ') + n, **blue}]
    if kind == 'Algorithm':
        return [{'t': f'Algorithm {n}', **blue}]
    start = (not sofar.strip()) or sofar.rstrip().endswith(('.', '?', '!', '“'))
    if supp:
        return [{'t': f'Supplementary Fig. {n}', **blue}]
    return [{'t': ('Figure ' if start else 'Fig. ') + n, **blue}]


def add_runs(par, runs, reg, size=None, font=None, base_bold=False, base_italic=False, mono_size=None):
    sofar = ''
    for r in runs:
        if 'br' in r:
            par.add_run().add_break()
            continue
        if 'cite' in r:
            run_ = par.add_run(cite_text(r['cite'], reg))
            big = (size or 9) >= 8
            set_run_font(run_, font, 7 if big else 5.5, color=BLUE, raise_pt=3 if big else 2.2)
            continue
        if 'xref' in r:
            for xr in xref_runs(r['xref'][0], r['xref'][1], sofar, reg):
                run_ = par.add_run(xr['t'])
                set_run_font(run_, font, size, bold=base_bold or None, italic=base_italic or None, color=BLUE if xr.get('blue') else None)
                sofar += xr['t']
            continue
        t = r['t']
        run_ = par.add_run(t)
        sofar += t
        name = font
        sz = size
        if r.get('mono'):
            name = MONO_FONT
            sz = (mono_size or (size - 1 if size else 8))
        set_run_font(run_, name, sz, bold=(True if (r.get('b') or base_bold) else None), italic=(True if (r.get('i') or base_italic) else None))
        if r.get('sup'):
            run_.font.superscript = True
        if r.get('sub'):
            run_.font.subscript = True
        if r.get('blue'):
            run_.font.color.rgb = BLUE
    return sofar


# --------------------------------------------------------------------------------------------------------- blocks
def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


def table_borders(tbl):
    tblpr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        for a, v in (('w:val', 'single'), ('w:sz', '2'), ('w:space', '0'), ('w:color', '000000')):
            e.set(qn(a), v)
        b.append(e)
    tblpr.append(b)
    mar = OxmlElement('w:tblCellMar')
    for side, w in (('top', 20), ('left', 60), ('bottom', 20), ('right', 60)):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(w)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tblpr.append(mar)


CH = 3.4      # average advance of a character in the 7 pt table face, in points
PAD = 7.0     # left and right cell margins


def col_widths(t):
    """Column widths in points and whether the table needs the full page width. A table is as wide as the text block
    when it fits there, the full page width otherwise, and no column is narrower than its longest word."""
    cols = t['cols']
    n = len(cols)
    nat, need = [0.0] * n, [0.0] * n
    for ri, row in enumerate(t['rows']):
        hdr = ri < t['header_rows']
        for c in row:
            if c['span'] != 1:
                continue
            txt = text_of(c['runs'])
            full = max((len(s) for s in txt.split('\n')), default=0)
            word = max((len(w) for w in re.split(r'\s+', txt) if w), default=1)
            nat[c['col']] = max(nat[c['col']], min(full, 44) * CH)
            need[c['col']] = max(need[c['col']], word * CH * (1.12 if hdr else 1.0))
    for i, c in enumerate(cols):
        if c['w']:
            nat[i] = max(nat[i], c['w'] * 28.35)
    floor = [max(nat[i] * 0.0, need[i]) + PAD for i in range(n)]
    widths = [max(nat[i], need[i]) + PAD for i in range(n)]
    if sum(widths) <= TEXT_W:
        return [w * TEXT_W / sum(widths) for w in widths], False
    wide = True
    while sum(widths) > WIDE_W:
        flex = [max(0.0, w - f) for w, f in zip(widths, floor)]
        if sum(flex) < 1e-6:
            widths = [w * WIDE_W / sum(widths) for w in widths]
            break
        k = min(1.0, (sum(widths) - WIDE_W) / sum(flex))
        widths = [w - fl * k for w, fl in zip(widths, flex)]
    return [w * WIDE_W / sum(widths) for w in widths], wide


def est_height(t, widths):
    """Rough rendered height of the table in points (used to decide whether it may be split across pages)."""
    h = 0.0
    for row in t['rows']:
        lines = 1
        for c in row:
            w = sum(widths[c['col']:c['col'] + c['span']]) - PAD
            n = 0
            for s in text_of(c['runs']).split('\n'):
                n += max(1, -(-len(s) * CH // max(w, 1)))
            lines = max(lines, n)
        h += lines * 8.6 + 3
    return h


def bookmark(par, name, uid):
    """A hidden bookmark around a paragraph, so that the layout probe can find where each float lands."""
    s = OxmlElement('w:bookmarkStart'); s.set(qn('w:id'), str(uid)); s.set(qn('w:name'), name)
    e = OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'), str(uid))
    ppr = par._p.find(qn('w:pPr'))
    if ppr is not None:
        ppr.addnext(s)
    else:
        par._p.insert(0, s)
    par._p.append(e)


def bm_name(label):
    return '_fl_' + re.sub(r'\W', '_', label or 'nolabel')


def add_table(doc, t, reg, uid=0):
    widths, wide = col_widths(t)
    t['wide'] = wide
    ncols = len(t['cols'])
    nrows = len(t['rows'])
    tbl = doc.add_table(rows=nrows, cols=ncols)
    tbl.autofit = False
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = tbl._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblpr.append(lay)
    ind = OxmlElement('w:tblInd'); ind.set(qn('w:w'), str(int((WIDE_INDENT if t['wide'] else 0) * 20))); ind.set(qn('w:type'), 'dxa'); tblpr.append(ind)
    tw = tblpr.find(qn('w:tblW'))
    if tw is None:
        tw = OxmlElement('w:tblW'); tblpr.append(tw)
    tw.set(qn('w:w'), str(int(sum(widths) * 20))); tw.set(qn('w:type'), 'dxa')
    table_borders(tbl)
    grid = tbl._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w * 20)))
    keep = est_height(t, widths) <= 190
    for ri, row in enumerate(t['rows']):
        tr = tbl.rows[ri]
        trpr = tr._tr.get_or_add_trPr()
        cs = OxmlElement('w:cantSplit'); trpr.append(cs)
        if ri < t['header_rows']:
            th = OxmlElement('w:tblHeader'); trpr.append(th)
        for c in row:
            j = c['col']
            span = c['span']
            tc = tbl.cell(ri, j)
            if span > 1:
                tc = tc.merge(tbl.cell(ri, j + span - 1))
            tcpr = tc._tc.get_or_add_tcPr()
            w_ = sum(widths[j:j + span])
            tcw = tcpr.find(qn('w:tcW'))
            if tcw is None:
                tcw = OxmlElement('w:tcW'); tcpr.insert(0, tcw)
            tcw.set(qn('w:w'), str(int(w_ * 20))); tcw.set(qn('w:type'), 'dxa')
            if c['vspan'] > 1:
                vm = OxmlElement('w:vMerge'); vm.set(qn('w:val'), 'restart'); tcpr.append(vm)
            elif c['vcont']:
                vm = OxmlElement('w:vMerge'); tcpr.append(vm)
            va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center' if (ri < t['header_rows'] or c['vspan'] > 1) else 'top'); tcpr.append(va)
            hdr = ri < t['header_rows']
            if hdr:
                shade(tc, HEAD_FILL)
            par = tc.paragraphs[0]
            par.style = doc.styles['SR Table text']
            colspec = t['cols'][j]
            if hdr:
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 and span == 1) or colspec['a'] == 'l' and span == 1 else WD_ALIGN_PARAGRAPH.CENTER
            else:
                par.alignment = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'r': WD_ALIGN_PARAGRAPH.RIGHT, 'c': WD_ALIGN_PARAGRAPH.CENTER}[colspec['a']] if span == 1 else WD_ALIGN_PARAGRAPH.CENTER
            if span > 1 and not hdr:
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if not c['vcont']:
                add_runs(par, c['runs'], reg, size=7, font=BODY_FONT, base_bold=hdr, mono_size=6.3)
            # short tables stay in one piece with their caption; a tall table may break between rows, but never
            # right after the header and never before its caption
            if keep or ri < t['header_rows'] + 1 or ri == nrows - 1:
                par.paragraph_format.keep_with_next = True
    bookmark(tbl.cell(0, 0).paragraphs[0], bm_name(t['label']), uid)
    # the caption sits below the table
    cap = doc.add_paragraph(style='SR Caption')
    return cap


def caption(doc, par, kind_label, runs, reg):
    r1 = par.add_run(kind_label + '.')
    set_run_font(r1, BODY_FONT, 9, bold=True)
    par.add_run('  ')
    add_runs(par, runs, reg, size=9, font=BODY_FONT)


def add_figure(doc, f, reg, label_text, uid=0):
    img = f['img']
    base = os.path.splitext(os.path.basename(img))[0]
    from PIL import Image
    w_px, h_px = Image.open(img).size
    if f['wide']:
        width = FIG_WIDTH.get(base, WIDE_W)
    else:
        width = FIG_WIDTH.get(base, FIG_DEFAULT)
    p = doc.add_paragraph(style='SR Figure')
    if f['wide']:
        p.paragraph_format.left_indent = Pt(WIDE_INDENT)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run().add_picture(img, width=Pt(width))
    bookmark(p, bm_name(f['label']), uid)
    cap = doc.add_paragraph(style='SR Caption')
    caption(doc, cap, label_text, f['caption'], reg)


def frame_affiliation(doc, reg):
    p = doc.add_paragraph(style='SR Affil')
    ppr = p._p.get_or_add_pPr()
    fr = OxmlElement('w:framePr')
    for a, v in (('w:w', str(int(TEXT_W * 20))), ('w:wrap', 'notBeside'), ('w:vAnchor', 'page'), ('w:hAnchor', 'page'),
                 ('w:x', str(int(LEFT * 20))), ('w:y', str(int(703 * 20)))):
        fr.set(qn(a), v)
    ppr.insert(0, fr)
    r1 = p.add_run('1'); set_run_font(r1, HEAD_FONT, 5.95, raise_pt=2.5)
    r2 = p.add_run(AFFIL); set_run_font(r2, HEAD_FONT, 8.5)
    r2.add_break()
    r3 = p.add_run(); set_run_font(r3, 'Wingdings', 5.95, raise_pt=2.5)
    sym = OxmlElement('w:sym'); sym.set(qn('w:font'), 'Wingdings'); sym.set(qn('w:char'), 'F02A'); r3._r.append(sym)
    r4 = p.add_run(' email: ' + EMAILS); set_run_font(r4, HEAD_FONT, 8.5)


def emit(doc_blocks, reg):
    d = Document()
    setup(d)
    # remove the empty first paragraph of the default template
    body = d.element.body
    for p in list(body.findall(qn('w:p'))):
        body.remove(p)

    # --- front matter ----------------------------------------------------------------------------------------
    t = d.add_paragraph(style='SR Title')
    t.paragraph_format.space_before = Pt(74)
    set_run_font(t.add_run(TITLE), HEAD_FONT, 26, bold=True)
    a = d.add_paragraph(style='SR Authors')
    for i, name in enumerate(AUTHORS):
        r_ = a.add_run(name); set_run_font(r_, HEAD_FONT, 10, bold=True)
        s_ = a.add_run('1'); set_run_font(s_, HEAD_FONT, 7, bold=True, raise_pt=3)
        w_ = a.add_run(); set_run_font(w_, 'Wingdings', 7, raise_pt=3)
        sym = OxmlElement('w:sym'); sym.set(qn('w:font'), 'Wingdings'); sym.set(qn('w:char'), 'F02A'); w_._r.append(sym)
        if i < len(AUTHORS) - 1:
            sep = a.add_run(' & ' if i == len(AUTHORS) - 2 else ', '); set_run_font(sep, HEAD_FONT, 10, bold=True)
    abstract = tex_abstract()
    ap = d.add_paragraph(style='SR Abstract')
    add_runs(ap, abstract, reg, size=9, font=HEAD_FONT, base_bold=True)
    kp = d.add_paragraph(style='SR Keywords')
    set_run_font(kp.add_run('Keywords'), HEAD_FONT, 10, bold=True)
    set_run_font(kp.add_run('  ' + KEYWORDS), BODY_FONT, 9)
    frame_affiliation(d, reg)

    # --- body ---------------------------------------------------------------------------------------------------
    pending_lead = None
    first_after_heading = True
    supp = False
    counters = {'Table': 0, 'Figure': 0, 'STable': 0, 'SFigure': 0}
    bm_uid = 100
    for b in doc_blocks:
        k = b[0]
        if k == 'marker':
            if b[1] == 'references':
                add_references(d, reg)
            if b[1] == 'supp':
                supp = True
            continue
        if k in ('h1', 'h2'):
            st = 'SR Heading 1' if k == 'h1' else 'SR Heading 2'
            hp = d.add_paragraph(style=st)
            add_runs(hp, b[1], reg, size=11 if k == 'h1' else 10, font=HEAD_FONT, base_bold=True)
            first_after_heading = True
            pending_lead = None
            continue
        if k == 'h3':
            pending_lead = b[1]
            continue
        if k == 'p':
            par = d.add_paragraph(style='SR Body first' if first_after_heading else 'Normal')
            if pending_lead is not None:
                lead = list(pending_lead)
                for r_ in lead:
                    r_['i'] = True
                lead.append({'t': '. '})
                add_runs(par, lead, reg, size=9, font=BODY_FONT)
                pending_lead = None
            add_runs(par, b[1], reg, size=9, font=BODY_FONT)
            first_after_heading = False
        elif k == 'list':
            if pending_lead is not None:
                par = d.add_paragraph(style='SR Body first')
                lead = [dict(r_, i=True) for r_ in pending_lead]
                add_runs(par, lead, reg, size=9, font=BODY_FONT)
                pending_lead = None
            for idx, it in enumerate(b[1]):
                par = d.add_paragraph(style='SR Bullet')
                if idx == 0:
                    par.paragraph_format.space_before = Pt(4)
                if idx == len(b[1]) - 1:
                    par.paragraph_format.space_after = Pt(4)
                par.add_run('•\t')
                tabs = OxmlElement('w:tabs'); tb = OxmlElement('w:tab'); tb.set(qn('w:val'), 'left'); tb.set(qn('w:pos'), str(int(9.9 * 20))); tabs.append(tb)
                par._p.get_or_add_pPr().append(tabs)
                add_runs(par, it, reg, size=9, font=BODY_FONT)
            first_after_heading = False
        elif k == 'quote':
            for ln in b[1]:
                par = d.add_paragraph(style='SR Quote')
                add_runs(par, ln, reg, size=7.5, font=MONO_FONT, mono_size=7.5)
            first_after_heading = False
        elif k == 'table':
            tt = b[1]
            key = 'STable' if supp else 'Table'
            counters[key] += 1
            lab = ('Supplementary Table S%d' if supp else 'Table %d') % counters[key]
            bm_uid += 1
            cap = add_table(d, tt, reg, bm_uid)
            caption(d, cap, lab, tt['caption'], reg)
            first_after_heading = False
        elif k == 'figure':
            ff = b[1]
            key = 'SFigure' if supp else 'Figure'
            counters[key] += 1
            lab = ('Supplementary Fig. S%d' if supp else 'Fig. %d') % counters[key]
            bm_uid += 1
            add_figure(d, ff, reg, lab, bm_uid)
            first_after_heading = False
        elif k == 'algorithm':
            al = b[1]
            for li, ln in enumerate(al['lines']):
                par = d.add_paragraph(style='SR Algorithm')
                if li == 0:
                    bm_uid += 1
                    bookmark(par, bm_name(al['label']), bm_uid)
                indent = len(ln) - len(ln.lstrip(' '))
                par.paragraph_format.left_indent = Pt(22 + indent * 3.2)
                bold = ln.startswith(('BEGIN', 'END', 'Step'))
                rr = par.add_run(ln.strip())
                set_run_font(rr, 'Cambria', 8.2, bold=bold or None, italic=True)
                par.paragraph_format.keep_with_next = True
            cap = d.add_paragraph(style='SR Caption')
            r1 = cap.add_run('Algorithm 1:'); set_run_font(r1, BODY_FONT, 9, bold=True)
            cap.add_run(' ')
            add_runs(cap, al['caption'], reg, size=9, font=BODY_FONT)
            first_after_heading = False
    d.core_properties.title = TITLE
    d.core_properties.author = '; '.join(AUTHORS)
    d.core_properties.subject = 'Manuscript'
    d.core_properties.keywords = KEYWORDS
    return d


def tex_abstract():
    src = open(os.path.join(T.PAPER, 'sections', 'abstract.tex'), encoding='utf8').read()
    blocks = T.parse_blocks(src)
    ab = [b for b in blocks if b[0] == 'abstract'][0]
    return ab[1]


def add_references(d, reg):
    h = d.add_paragraph(style='SR Heading 1')
    add_runs(h, [T.run('References')], reg, size=11, font=HEAD_FONT, base_bold=True)
    for n, key in enumerate(reg.cite_order, 1):
        par = d.add_paragraph(style='SR Reference')
        tabs = OxmlElement('w:tabs'); tb = OxmlElement('w:tab'); tb.set(qn('w:val'), 'left'); tb.set(qn('w:pos'), str(int(12.8 * 20))); tabs.append(tb)
        par._p.get_or_add_pPr().append(tabs)
        r_ = par.add_run(f'{n}.\t'); set_run_font(r_, BODY_FONT, 7.5)
        add_runs(par, [dict(x) for x in REFS[key]], reg, size=7.5, font=BODY_FONT)


ORDER = {
    'pPr': ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs',
            'suppressAutoHyphens', 'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
            'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection', 'textAlignment',
            'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange'],
    'rPr': ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps', 'strike', 'dstrike', 'outline', 'shadow', 'emboss', 'imprint', 'noProof',
            'snapToGrid', 'vanish', 'webHidden', 'color', 'spacing', 'w', 'kern', 'position', 'sz', 'szCs', 'highlight', 'u', 'effect', 'bdr', 'shd',
            'fitText', 'vertAlign', 'rtl', 'cs', 'em', 'lang', 'eastAsianLayout', 'specVanish', 'oMath'],
    'tblPr': ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd',
              'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription'],
    'tcPr': ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark'],
}


def normalize_order(root):
    """python-docx and the direct XML edits above append properties in the order they were made; the schema wants a fixed
    order, which Word tolerates but stricter readers do not."""
    W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    for tag, order in ORDER.items():
        rank = {W + n: i for i, n in enumerate(order)}
        for el in root.iter(W + tag):
            kids = list(el)
            kids.sort(key=lambda k: rank.get(k.tag, 999))
            for k in kids:
                el.remove(k)
            for k in kids:
                el.append(k)


def apply_delays(doc, delays):
    """Moves a float past the given number of following paragraphs (never across the references or supplementary markers)."""
    for label, n in delays.items():
        idx = next((i for i, b in enumerate(doc) if b[0] in ('table', 'figure', 'algorithm') and b[1].get('label') == label), None)
        if idx is None:
            continue
        b = doc.pop(idx)
        i, hops = idx, 0
        # a float never passes another float of its own kind, so the numbers still follow the order of first mention
        while hops < n and i < len(doc) and doc[i][0] not in ('marker', 'h1', b[0]):
            if doc[i][0] in ('p', 'list'):
                hops += 1
            i += 1
        # a negative delay moves the float back past that many paragraphs (not across a heading)
        while hops < -n and i > 0 and doc[i - 1][0] not in ('marker', 'h1', 'h2', b[0]):
            i -= 1
            if doc[i][0] in ('p', 'list'):
                hops += 1
        while i > 0 and doc[i - 1][0] == 'h3':      # keep a run-in lead with its paragraph
            i -= 1
        if 0 < i and doc[i - 1][0] in ('h1', 'h2'):   # never directly under a heading
            i = min(len(doc), i + 1)
        doc.insert(i, b)
    return doc


def main():
    blocks = assemble()
    dpath = os.path.join(HERE, 'float_delays.json')
    if os.path.exists(dpath):
        import json
        blocks = apply_delays(blocks, json.load(open(dpath)))
    # the same flow order is used for numbering and for emission
    reg = number(blocks)
    d = emit(blocks, reg)
    normalize_order(d.element)
    normalize_order(d.styles.element)
    for sec in d.sections:
        for hf in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
            normalize_order(hf._element)
    d.save(OUT)
    print('saved', OUT)
    print('citations:', len(reg.cite_order), '| tables:', sum(1 for b in blocks if b[0] == 'table'), '| figures:', sum(1 for b in blocks if b[0] == 'figure'))
    unused = set(REFS) - set(reg.cite_order)
    print('references not cited:', sorted(unused))
    missing = [k for k in reg.cite_order if k not in REFS]
    print('cited but missing from refs_data:', missing)


if __name__ == '__main__':
    main()
